"""Desktop acceptance reads must not accept unrecalculated or edited artifacts."""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pytest

pytest.importorskip("xlsxwriter")
openpyxl = pytest.importorskip("openpyxl")

spec = importlib.util.spec_from_file_location(
    "excel_parity_tool", Path(__file__).parents[1] / "tools/check_excel_parity.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
check_edits = module.check_edits


def test_desktop_edit_generate_does_not_claim_recalculation(tmp_path):
    generated = check_edits(tmp_path, "generate")
    assert generated["status"] == "NOT_RECALCULATED"
    assert generated["cases"] == 2
    assert check_edits(tmp_path, "read")["status"] == "MISMATCH"


def test_desktop_edit_original_identity_is_checked(tmp_path):
    check_edits(tmp_path, "generate")
    path = tmp_path / "value-identity.json"
    identity = json.loads(path.read_text())
    identity["source_sha256"] = "wrong"
    path.write_text(json.dumps(identity))
    with pytest.raises(ValueError, match="source identity changed"):
        check_edits(tmp_path, "read")


@pytest.mark.parametrize("sheet,cell,value", [
    ("Results0", "B1", "=2"),
    ("Inputs0", "B2", 999),
    ("Results0", "C1", 123),
])
def test_desktop_edit_detects_unapproved_formula_input_or_reference_change(tmp_path, sheet, cell, value):
    check_edits(tmp_path, "generate")
    path = tmp_path / "value.xlsx"
    book = openpyxl.load_workbook(path)
    book[sheet][cell] = value
    book.save(path)
    book.close()
    report = check_edits(tmp_path, "read")
    assert report["status"] == "MISMATCH"
    assert any(f.get("sheet") == sheet and f.get("cell") == cell for f in report["failures"])


@pytest.mark.parametrize("extra", [None, "=123", 999, "unexpected"])
def test_desktop_identity_allows_only_empty_used_range_padding(tmp_path, extra):
    from openpyxl.styles import PatternFill

    check_edits(tmp_path, "generate")
    source = tmp_path / "value-edited-source.xlsx"
    path = tmp_path / "value.xlsx"
    book = openpyxl.load_workbook(path)
    book["Readme"]["Z20"].fill = PatternFill("solid", fgColor="FF0000")
    book["Readme"]["Z20"] = extra
    book.save(path)
    book.close()
    failures = module.inspect_identity(source, path)
    if extra is None:
        assert failures == []
    else:
        assert any(f.get("cell") == "Z20" for f in failures)
