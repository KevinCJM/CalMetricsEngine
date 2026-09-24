"""Generate real formula conformance workbooks, recalculate, inspect and report.

Native Excel acceptance: Windows + installed Excel + pywin32, dedicated COM instance.
LibreOffice is an independent supplementary lane, never labeled Microsoft Excel.
``--engine read`` reads files already explicitly recalculated in desktop Excel.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import re
import shutil
import subprocess
import tempfile
import time
from decimal import Decimal
from itertools import zip_longest
from pathlib import Path

import openpyxl
import xlsxwriter

from calmetrics_engine import build_info


def write_bundle(cases, destination):
    """Serialization-only test fixture packer; all formula references are native."""
    manifest = []
    total = 0
    with tempfile.TemporaryDirectory(prefix="calmetrics-excel-cases-") as directory:
        with xlsxwriter.Workbook(destination, {"constant_memory": True, "strings_to_formulas": False,
                                             "strings_to_urls": False, "tmpdir": directory}) as book:
            book.set_calc_mode("auto")
            for case_id, (name, make_plan) in enumerate(cases):
                plan = make_plan()
                metadata = plan.metadata()
                total += metadata["cells"]
                if total > 2_000_000:
                    raise RuntimeError("OVER_BUDGET: conformance workbook")
                reference = plan.reference()
                prefix = f"C{case_id:03}_"
                stream = Path(directory) / "cells.jsonl"
                plan.write_cells(str(stream), prefix)
                sheets = {}
                for sheet in metadata["sheets"]:
                    target = book.add_worksheet(prefix + sheet)
                    target.set_column(0, 0, 32)
                    target.set_column(1, 4, 24)
                    sheets[prefix + sheet] = target
                with stream.open() as source:
                    for line in source:
                        sheet, row, column, kind, value = json.loads(line)
                        if kind == "reference":
                            root, index = map(int, value.split(":"))
                            r = reference[root]
                            if r.get("statuses") and r["statuses"][index]:
                                kind, value = "text", "ERROR:STATUS:" + str(r["statuses"][index])
                            elif r["error"]:
                                kind, value = "text", "ERROR:" + r["error"]
                            else:
                                value = r["values"][index]
                                dtype = metadata["outputs"][root]["dtype"]
                                if dtype == "bool":
                                    kind = "boolean"
                                elif isinstance(value, float) and not math.isfinite(value):
                                    kind, value = "text", "NaN" if math.isnan(value) else "+Inf" if value > 0 else "-Inf"
                                else:
                                    kind = "number"
                        target = sheets[sheet]
                        if kind == "formula":
                            target.write_formula(row, column, "=" + value, None, "NOT_RECALCULATED")
                        elif kind == "number":
                            target.write_number(row, column, float(value))
                        elif kind == "boolean":
                            target.write_boolean(row, column, bool(int(value)))
                        else:
                            target.write_string(row, column, str(value))
                manifest.append({"case": name, "prefix": prefix, "plan": metadata,
                                 "reference": reference})
    return manifest


def inspect_identity(source_path, calculated_path, *, skip_reference=False):
    failures = []
    source = openpyxl.load_workbook(source_path, data_only=False, read_only=True)
    formulas = openpyxl.load_workbook(calculated_path, data_only=False, read_only=True)
    def normalized(value):
        from openpyxl.formula import Tokenizer
        result = []
        tokens = iter(Tokenizer(value).items)
        for token in tokens:
            text = token.value
            if token.type == "FUNC":
                text = text.replace("_xlfn.", "").upper()
                if text in {"TRUE(", "FALSE("}:
                    closing = next(tokens, None)
                    if closing is None or closing.value != ")":
                        raise ValueError("invalid boolean formula")
                    text = text[:-1]
            elif token.type == "OPERAND" and token.subtype == "RANGE":
                text = text.replace("$", "")
                text = re.sub(r"'([A-Za-z0-9_]+)'!", r"\1!", text)
            elif token.type == "OPERAND" and token.subtype == "NUMBER":
                text = str(Decimal(text).normalize())
            if token.type != "WHITE-SPACE":
                result.append(text)
        return "".join(result)
    for expected_sheet in source:
        if expected_sheet.title not in formulas:
            failures.append({"identity_error": "missing sheet", "sheet": expected_sheet.title})
            continue
        actual_sheet = formulas[expected_sheet.title]
        # Excel can extend the used range to formatted, empty columns/rows.
        # Compare the union: allow formatting-only padding, never added content.
        for row, (expected_row, actual_row) in enumerate(zip_longest(expected_sheet, actual_sheet, fillvalue=()), 1):
            for column, (expected, actual) in enumerate(zip_longest(expected_row, actual_row), 1):
                before, after = getattr(expected, "value", None), getattr(actual, "value", None)
                before_type, after_type = getattr(expected, "data_type", None), getattr(actual, "data_type", None)
                if before_type == "f":
                    same = after_type == "f" and normalized(before) == normalized(after)
                elif isinstance(before, bool) and after_type == "f":
                    same = normalized(after) == str(before).upper()
                elif skip_reference and "Results" in expected_sheet.title and column == 3:
                    continue  # Compared independently with the manifest below.
                else:
                    same = before == after or before == "" and after is None
                if not same:
                    failures.append({"identity_error": "formula/input changed", "sheet": expected_sheet.title,
                                     "cell": openpyxl.utils.get_column_letter(column) + str(row)})
    for title in set(formulas.sheetnames) - set(source.sheetnames):
        failures.append({"identity_error": "unexpected sheet", "sheet": title})
    source.close()
    formulas.close()
    return failures


def inspect(path, manifest):
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    failures = inspect_identity(path.with_name("source.xlsx"), path, skip_reference=True)
    checked = 0
    for case in manifest:
        offset = 0
        for output, reference in zip(case["plan"]["outputs"], case["reference"], strict=True):
            for i in range(len(output["addresses"])):
                title = case["prefix"] + "Results" + str(offset // 1_000_000)
                row = offset % 1_000_000 + 1
                offset += 1
                checked += 1
                if title not in book:
                    failures.append({"case": case["case"], "missing_sheet": title})
                    continue
                values = next(book[title].iter_rows(min_row=row, max_row=row, values_only=True))
                status = reference.get("statuses", [])
                if status and status[i]:
                    expected = "ERROR:STATUS:" + str(status[i])
                elif reference["error"]:
                    expected = "ERROR:" + reference["error"]
                else:
                    expected = reference["values"][i]
                    if output["dtype"] == "bool":
                        expected = bool(expected)
                    elif isinstance(expected, float) and not math.isfinite(expected):
                        expected = "NaN" if math.isnan(expected) else "+Inf" if expected > 0 else "-Inf"
                def matches(actual, expected=expected, output=output, case=case):
                    if isinstance(expected, str):
                        return actual == expected
                    if output["dtype"] == "bool":
                        return isinstance(actual, bool) and actual == expected
                    if not isinstance(actual, int | float) or isinstance(actual, bool):
                        return False
                    if output["dtype"] == "int64":
                        return actual == expected
                    return abs(actual - expected) <= case["plan"]["atol"] + case["plan"]["rtol"] * abs(expected)
                # Compare independently with frozen C++ values; a worksheet's
                # editable MATCH label is never sufficient evidence.
                if len(values) < 5 or values[4] != "MATCH" or not matches(values[1]) or not matches(values[2]):
                    failures.append({"case": case["case"], "sheet": title, "row": row,
                                     "excel": values[1], "cpp": expected, "comparison": values[4]})
    book.close()
    return {"status": "PASS" if not failures and checked else "MISMATCH", "checked_values": checked,
            "failures": failures}


def recalculate(workbook, engine):
    """Force a real application calculation; never synthesize cached values."""
    if engine == "excel":
        import win32com.client
        application = win32com.client.DispatchEx("Excel.Application")
        application.DisplayAlerts = False
        application.AutomationSecurity = 3
        book = None
        try:
            engine_version = application.Version + "/" + str(application.Build)
            book = application.Workbooks.Open(str(workbook.absolute()), UpdateLinks=0)
            application.CalculateFullRebuild()
            book.Save()
        finally:
            if book is not None:
                book.Close(SaveChanges=False)
            application.Quit()
    elif engine == "libreoffice":
        binary = shutil.which("libreoffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        engine_version = subprocess.check_output([binary, "--version"], text=True).strip()
        with tempfile.TemporaryDirectory(prefix="calmetrics-lo-") as directory:
            target = Path(directory) / "converted"
            target.mkdir()
            profile = Path(directory) / "profile"
            (profile / "user").mkdir(parents=True)
            (profile / "user/registrymodifications.xcu").write_text(
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
                '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>'
                '</item></oor:items>')
            subprocess.run([binary, "-env:UserInstallation=" + (Path(directory) / "profile").as_uri(),
                            "--headless", "--convert-to", "xlsx", "--outdir", str(target), str(workbook)],
                           check=True, timeout=180, capture_output=True, text=True)
            shutil.copy2(target / workbook.name, workbook)
    return engine_version


def check_edits(directory, engine):
    """Check live dependencies and stale references using deliberately edited inputs."""
    import numpy as np

    from calmetrics_engine import GraphCompiler, excel, operators

    directory.mkdir(parents=True, exist_ok=True)
    x = np.array([1., 2., 3.])
    mean_graph = GraphCompiler({"x": "series"}).compile("mean(x)")
    rolling_graph = GraphCompiler({"x": "series", "width": "scalar"}).compile(
        "rolling_mean(x,width)", result_format="typed")
    cases = [
        ("value", excel.plan(mean_graph, {"x": x}), "B1", 9.),
        ("layout", excel.plan(rolling_graph, {"x": x}, {"width": 2.}), "B4", 3.),
    ]
    failures = []
    for name, plan, address, changed in cases:
        workbook = directory / (name + ".xlsx")
        identity_path = directory / (name + "-identity.json")
        if engine != "read":
            excel.export(plan, workbook)
            shutil.copy2(workbook, directory / (name + "-original.xlsx"))
            book = openpyxl.load_workbook(workbook)
            book["Inputs0"][address] = changed
            book.save(workbook)
            book.close()
            # A separate frozen artifact records the intentional edit. Desktop
            # readback must preserve every formula and every other input.
            edited_source = directory / (name + "-edited-source.xlsx")
            shutil.copy2(workbook, edited_source)
            identity_path.write_text(json.dumps({
                "plan_identity": plan.identity,
                "source_sha256": hashlib.sha256(edited_source.read_bytes()).hexdigest(),
            }))
        else:
            identity = json.loads(identity_path.read_text())
            edited_source = directory / (name + "-edited-source.xlsx")
            if identity["plan_identity"] != plan.identity:
                raise ValueError("input edit fixture build/plan changed")
            if hashlib.sha256(edited_source.read_bytes()).hexdigest() != identity["source_sha256"]:
                raise ValueError("edited fixture source identity changed")
        if engine == "generate":
            continue
        if engine in {"excel", "libreoffice"}:
            recalculate(workbook, engine)
        failures.extend(inspect_identity(edited_source, workbook))
        book = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
        result = list(book["Results0"].values)
        if name == "value":
            expected = operators.mean(np.array([9., 2., 3.]))
            valid = (isinstance(result[0][1], int | float)
                     and abs(result[0][1] - expected) <= 1e-12 + 1e-10 * abs(expected)
                     and result[0][2] == 2.)
        else:
            valid = all(row[1] == "LAYOUT_STALE: re-export required" for row in result)
        valid = valid and all(row[4] == "REFERENCE_STALE" for row in result)
        valid = valid and book["Readme"]["B7"].value == "REFERENCE_STALE"
        if not valid:
            failures.append({"case": name, "values": result})
        book.close()
    return {"status": "NOT_RECALCULATED" if engine == "generate" else "PASS" if not failures else "MISMATCH",
            "cases": len(cases), "failures": failures}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--engine", choices=["generate", "read", "excel", "libreoffice"], required=True)
    parser.add_argument("--all-cases", action="store_true")
    parser.add_argument("--check-edits", action="store_true",
                        help="Generate/recalculate/read value and layout edit fixtures with this engine")
    parser.add_argument("--filter", default="")
    parser.add_argument("--engine-version", default="external Excel; version not supplied")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    workbook = args.output / "conformance.xlsx"
    manifest_path = args.output / "manifest.json"
    start = time.perf_counter()
    if args.engine != "read":
        source = Path(__file__).parents[1] / "tests/excel_cases.py"
        spec = importlib.util.spec_from_file_location("excel_cases", source)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        selected = [c for c in module.cases(args.all_cases) if not args.filter or args.filter in c[0]]
        manifest = write_bundle(selected, workbook)
        manifest_path.write_text(json.dumps({"build": build_info(), "cases": manifest,
                                            "generated_sha256": hashlib.sha256(workbook.read_bytes()).hexdigest()}, indent=2))
        shutil.copy2(workbook, workbook.with_name("source.xlsx"))
    else:
        saved = json.loads(manifest_path.read_text())
        if hashlib.sha256(workbook.with_name("source.xlsx").read_bytes()).hexdigest() != saved["generated_sha256"]:
            raise ValueError("original formula artifact identity changed")
        manifest = saved["cases"]
    engine_version = args.engine_version if args.engine == "read" else "not recalculated"
    if args.engine in {"excel", "libreoffice"}:
        engine_version = recalculate(workbook, args.engine)
    result = {"status": "NOT_RECALCULATED"} if args.engine == "generate" else inspect(workbook, manifest)
    if args.check_edits:
        result["input_edits"] = check_edits(args.output / "input-edits", args.engine)
        if result["input_edits"]["status"] not in {"PASS", "NOT_RECALCULATED"}:
            result["status"] = "MISMATCH"
    result.update({"engine": args.engine, "engine_version": engine_version, "cases": len(manifest),
                   "elapsed_seconds": time.perf_counter() - start, "bytes": workbook.stat().st_size,
                   "workbook_sha256": hashlib.sha256(workbook.read_bytes()).hexdigest()})
    (args.output / "report.json").write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))
    return 0 if result["status"] in {"PASS", "NOT_RECALCULATED"} else 1


if __name__ == "__main__":
    raise SystemExit(main())
